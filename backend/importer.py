import re
from io import BytesIO
from openpyxl import load_workbook

def initial_password(name, roll):
    return re.sub(r'\s+', '', name) + re.sub(r'\s+', '', roll)

def parse_workbook(content):
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    headers = None
    result, errors = [], []
    seen = set()
    for line, values in enumerate(rows, 1):
        cells = [str(x).strip() if x is not None else '' for x in values]
        lowered = [x.lower() for x in cells]
        if headers is None:
            if 'email' in lowered and 'name' in lowered:
                headers = lowered
            continue
        if not any(cells):
            continue
        row = dict(zip(headers, cells))
        email = row.get('email', '').lower()
        name = row.get('name', '')
        roll_keys = [h for h in headers if 'roll' in h and 'student' not in h]
        members = [{'name': name, 'roll': row.get(roll_keys[0], '') if roll_keys else ''}]
        for n in [2,3,4]:
            roll_key = next((h for h in headers if 'roll' in h and f'student {n}' in h), '')
            members.append({'name': row.get(f'name of student {n}', ''), 'roll': row.get(roll_key, '')})
        if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', email):
            errors.append({'row':line, 'error':'Missing or invalid leader email'})
        elif any(not m['name'] or not m['roll'] for m in members):
            errors.append({'row':line, 'error':'All four member names and roll numbers are required'})
        elif email in seen:
            errors.append({'row':line, 'error':'Duplicate leader email in workbook'})
        elif len(set(m['roll'].lower() for m in members)) != 4:
            errors.append({'row':line, 'error':'Duplicate member roll number within team'})
        else:
            seen.add(email)
            result.append({'name': f"Team {name}", 'email':email, 'members':members,
                           'password':initial_password(name, members[0]['roll']), 'row':line})
    workbook.close()
    if headers is None:
        raise ValueError('Could not find the Name and Email header row in the first worksheet')
    return {'teams':result, 'errors':errors}
